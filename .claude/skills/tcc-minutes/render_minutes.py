#!/usr/bin/env python3
"""
Render TCC-format meeting minutes (.docx) from a JSON spec.

Usage:
    python3 render_minutes.py --input spec.json --output "Minutes.docx"

The JSON spec is data-driven so the same house style is applied to any meeting.
See spec.schema.md (next to this file) for the full shape. Minimal example:

{
  "subtitle": "Straight Bat Private Equity Pty Ltd | 9 June 2026",
  "subheading": "Governance Review — Implementation Kick-off",
  "status": "CONFIDENTIAL — DRAFT",
  "details": [["Entity", "Straight Bat Private Equity Pty Ltd"], ["Date", "9 June 2026"]],
  "attendees": [{"name": "Kim Dema", "role": "Company Secretary", "status": "Present"}],
  "sections": [
    {"heading": "1. Opening", "presenter": "Chair: ... | Time allocated: 5 minutes",
     "blocks": [{"type": "para", "text": "..."}, {"type": "resolution", "text": "..."}]}
  ],
  "actions": [{"ref": "A1", "action": "...", "owner": "Kim Dema", "due": "Next Meeting"}],
  "confirmation": {"statement": "...",
                   "signatories": [{"role": "Chair", "name": "Becky Barr", "title": "Compliance Manager"}]},
  "footnote": "Draft prepared from the Granola transcript ..."
}
"""
import argparse
import json
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x60, 0x60, 0x60)
RED = RGBColor(0xB0, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

DEFAULT_TITLE = "TRUSTEE COMMITTEE MEETING — MINUTES OF MEETING"


def build(spec: dict) -> Document:
    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10.5)

    def heading(text, level=1):
        h = doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.color.rgb = NAVY
        return h

    def para(text=None, bold=False, italic=False, size=None, color=None,
             align=None, space_after=6):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return p

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(text)
        p.paragraph_format.space_after = Pt(3)
        return p

    def kv_table(rows):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for k, v in rows:
            cells = t.add_row().cells
            cells[0].width = Inches(1.8)
            cells[1].width = Inches(4.7)
            rk = cells[0].paragraphs[0].add_run(str(k))
            rk.bold = True
            cells[1].paragraphs[0].add_run(str(v))
        return t

    def grid_table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        for i, h in enumerate(headers):
            run = hdr[i].paragraphs[0].add_run(h)
            run.bold = True
            run.font.color.rgb = WHITE
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].paragraphs[0].add_run(str(val))
                if widths:
                    cells[i].width = widths[i]
        return t

    # ---- Title block ----
    para(spec.get("title", DEFAULT_TITLE), bold=True, size=15, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if spec.get("subtitle"):
        para(spec["subtitle"], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if spec.get("subheading"):
        para(spec["subheading"], italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if spec.get("status"):
        para(spec["status"], bold=True, color=RED,
             align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # ---- Meeting details ----
    if spec.get("details"):
        heading("Meeting Details", 2)
        kv_table(spec["details"])

    # ---- Attendees ----
    if spec.get("attendees"):
        heading("Attendees", 2)
        grid_table(
            ["Name", "Title / Role", "Status"],
            [[a.get("name", ""), a.get("role", ""), a.get("status", "")]
             for a in spec["attendees"]],
            widths=[Inches(1.5), Inches(3.8), Inches(1.2)],
        )

    # ---- Sections ----
    for sec in spec.get("sections", []):
        heading(sec["heading"], 1)
        if sec.get("presenter"):
            para(sec["presenter"], italic=True, color=GREY, space_after=6)
        for block in sec.get("blocks", []):
            btype = block.get("type", "para")
            text = block.get("text", "")
            if btype == "sub":
                para(text, bold=True, space_after=3)
            elif btype == "bullet":
                bullet(text)
            elif btype == "resolution":
                prefix = block.get("prefix", "RESOLUTION: ")
                para(prefix + text, bold=True)
            else:  # para
                para(text)

    # ---- Action Register ----
    if spec.get("actions"):
        heading("Action Register", 1)
        grid_table(
            ["Ref", "Action Item", "Owner", "Due By"],
            [[a.get("ref", ""), a.get("action", ""), a.get("owner", ""), a.get("due", "")]
             for a in spec["actions"]],
            widths=[Inches(0.5), Inches(4.3), Inches(1.4), Inches(1.0)],
        )

    # ---- Confirmation ----
    conf = spec.get("confirmation")
    if conf:
        heading("Confirmation of Minutes", 1)
        if conf.get("statement"):
            para(conf["statement"])
        for sig in conf.get("signatories", []):
            para(sig.get("role", ""), bold=True, space_after=0)
            line = sig.get("name", "")
            if sig.get("title"):
                line = f"{line} | {sig['title']}"
            para(line, space_after=0)
            para("Date: ________________________", space_after=12)

    if spec.get("footnote"):
        para(spec["footnote"], italic=True, size=8.5, color=GREY)

    return doc


def main():
    ap = argparse.ArgumentParser(description="Render TCC-format minutes to .docx from a JSON spec.")
    ap.add_argument("--input", "-i", required=True, help="Path to the JSON spec file.")
    ap.add_argument("--output", "-o", required=True, help="Path to the output .docx file.")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as fh:
        spec = json.load(fh)

    doc = build(spec)
    doc.save(args.output)
    print(f"Saved: {args.output}")


if __name__ == "__main__":
    sys.exit(main())
